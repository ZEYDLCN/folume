"""A second, independent producer for the synthetic papers: DOCX typeset by LibreOffice.

The reportlab corpus is laid out by code that knows what it drew. LibreOffice lays the same content out with its
own justification, kerning, font substitution and two-column flow, which is a closer stand-in for PDFs from a real
word processor. Used by an optional test (skipped when LibreOffice is missing):

    python -m synth.docx_paper out_dir [name]
"""

from __future__ import annotations

import io
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw

from .corpus import SPECS
from .spec import Bullets, Code, Equation, Figure, Heading, Paper, Para, Table
from .textgen import build_paper

BODY_FONT = "FreeSerif"
MONO_FONT = "DejaVu Sans Mono"

_GREEK = {r"\theta": "θ", r"\lambda": "λ", r"\sum": "Σ", r"\alpha": "α", r"\beta": "β", r"\mu": "μ", r"\sigma": "σ",
          r"\nabla": "∇", r"\cdot": "·", r"\in": "∈", r"\|": "‖"}


def soffice_available() -> bool:
    return shutil.which("soffice") is not None


def _plain_math(expr: str) -> str:
    for k, v in _GREEK.items():
        expr = expr.replace(k, v)
    expr = re.sub(r"\\(?:frac|mathbb|mathcal|text|operatorname)\{([^}]*)\}\{([^}]*)\}", r"(\1)/(\2)", expr)
    expr = re.sub(r"\\[a-zA-Z]+", "", expr)
    return re.sub(r"[{}]", "", expr).replace("^", "").replace("_", "")


def _set_cols(section, n: int, space_cm: float = 0.7) -> None:
    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:cols")):
        sect_pr.remove(old)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(space_cm * 567)))
    sect_pr.append(cols)


def plain_text(text: str) -> str:
    """What a paragraph of the synthetic markup reads like once typeset (citations as [n], math as text)."""
    text = re.sub(r"\{cite:([\d,]+)\}", lambda m: "[" + ", ".join(m.group(1).split(",")) + "]", text)
    text = re.sub(r"\$([^$]*)\$", lambda m: _plain_math(m.group(1)), text)
    return re.sub(r"</?[ib]>", "", text)


def _add_runs(par, text: str) -> None:
    """Fills a paragraph from the synthetic markup (<i>, <b>, {cite:N}, $math$)."""
    text = re.sub(r"\{cite:([\d,]+)\}", lambda m: "[" + ", ".join(m.group(1).split(",")) + "]", text)
    text = re.sub(r"\$([^$]*)\$", lambda m: _plain_math(m.group(1)), text)
    for part in re.split(r"(<i>.*?</i>|<b>.*?</b>)", text):
        if not part:
            continue
        if part.startswith("<i>"):
            run = par.add_run(re.sub(r"</?i>", "", part))
            run.italic = True
        elif part.startswith("<b>"):
            run = par.add_run(re.sub(r"</?b>", "", part))
            run.bold = True
        else:
            par.add_run(re.sub(r"</?[ib]>", "", part))


def _picture(kind: str, wide: bool) -> io.BytesIO:
    w, h = (900, 420) if wide else (520, 360)
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([50, 20, w - 20, h - 50], outline="black", width=2)
    pts = [(50 + i * (w - 80) / 8, h - 60 - (i * i * 3 % 190) - 10) for i in range(9)]
    if kind == "bar":
        for i, (x, y) in enumerate(pts[:8]):
            d.rectangle([x + 4, y, x + 30, h - 52], fill=("#3366cc", "#cc6633")[i % 2])
    else:
        d.line(pts, fill="#3366cc", width=3)
        d.line([(x, y + 25) for x, y in pts], fill="#cc3333", width=3)
    d.text((w // 2 - 20, h - 40), "epoch", fill="black")
    out = io.BytesIO()
    img.save(out, format="PNG")
    out.seek(0)
    return out


def _page_number_field(par) -> None:
    run = par.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)


def build_docx(paper: Paper, path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
    head = sec.header.paragraphs[0]
    head.text = paper.short_title or paper.title[:50]
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    head.runs[0].italic = True
    head.runs[0].font.size = Pt(8)
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _page_number_field(foot)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run(paper.title)
    run.bold = True
    run.font.size = Pt(20)
    for name, aff in paper.authors:
        a = doc.add_paragraph()
        a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        a.add_run(name).bold = True
        a.add_run(f", {aff}").font.size = Pt(9)
    ab = doc.add_paragraph()
    ab.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ab.paragraph_format.left_indent = ab.paragraph_format.right_indent = Cm(1.0)
    ab.add_run("Abstract—").bold = True
    ab.add_run(paper.abstract)
    kw = doc.add_paragraph()
    kw.paragraph_format.left_indent = kw.paragraph_format.right_indent = Cm(1.0)
    kw.add_run("Keywords: ").italic = True
    kw.add_run(", ".join(paper.keywords))

    body_sec = doc.add_section(WD_SECTION.CONTINUOUS)
    _set_cols(body_sec, 2)
    n_head = n_fig = n_tab = n_eq = 0
    for item in paper.body:
        if isinstance(item, Heading):
            if item.level == 1:
                n_head += 1
                label = f"{n_head}  {item.text.upper()}"
                sub = 0
            else:
                sub += 1
                label = f"{n_head}.{sub}  {item.text}"
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(label)
            r.bold = True
            r.font.size = Pt(11 if item.level == 1 else 10)
        elif isinstance(item, Para):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(0.4)
            _add_runs(p, item.text)
        elif isinstance(item, Bullets):
            for i, it in enumerate(item.items, 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.first_line_indent = Cm(-0.4)
                _add_runs(p, f"{i}. {it}" if item.ordered else f"• {it}")
        elif isinstance(item, Equation):
            n_eq += 1
            p = doc.add_paragraph()
            width = doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin
            p.paragraph_format.tab_stops.add_tab_stop(int(width / 2 * 0.5), WD_TAB_ALIGNMENT.CENTER)
            p.paragraph_format.tab_stops.add_tab_stop(int(width / 2 - Cm(0.35)), WD_TAB_ALIGNMENT.RIGHT)
            r = p.add_run(f"\t{_plain_math(item.expr)}\t({n_eq})")
            r.italic = True
        elif isinstance(item, Figure):
            n_fig += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(_picture(item.kind, False), width=Cm(7.8))
            c = doc.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = c.add_run(f"Fig. {n_fig}. ")
            r.font.size = Pt(9)
            r2 = c.add_run(re.sub(r"<[^>]+>", "", item.caption))
            r2.font.size = Pt(9)
        elif isinstance(item, Table):
            n_tab += 1
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = cap.add_run(f"TABLE {n_tab}. {item.caption.upper()}")
            cr.font.size = Pt(8)
            table = doc.add_table(rows=1 + len(item.rows), cols=len(item.header))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            col_w = Cm(8.2) / len(item.header)  # fits one text column
            for col in table.columns:
                col.width = int(col_w)
                for cell in col.cells:
                    cell.width = int(col_w)
            _borders(table)
            for j, h in enumerate(item.header):
                table.rows[0].cells[j].text = h
                table.rows[0].cells[j].paragraphs[0].runs[0].bold = True
            for i, row in enumerate(item.rows, 1):
                for j, cell in enumerate(row):
                    table.rows[i].cells[j].text = str(cell)
            for row in table.rows:
                for cell in row.cells:
                    for par in cell.paragraphs:
                        for run in par.runs:
                            run.font.size = Pt(8)
            doc.add_paragraph()
        elif isinstance(item, Code):
            for line in item.lines:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(line)
                r.font.name = MONO_FONT
                r.font.size = Pt(8)

    if not any(isinstance(i, Heading) and i.text.lower().startswith("reference") for i in paper.body):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        hr = h.add_run("REFERENCES")
        hr.bold = True
        hr.font.size = Pt(11)
    for i, ref in enumerate(paper.references, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        names = ", ".join(f"{f[0]}. {last}" for f, last in ref.authors)
        pages = f", pp. {ref.pages}" if ref.pages else ""
        r = p.add_run(f"[{i}] {names}, \"{ref.title},\" {ref.venue}{pages}, {ref.year}.")
        r.font.size = Pt(8)
    doc.save(str(path))


def to_pdf(docx_path: Path, out_dir: Path) -> Path:
    """Converts with a throw-away LibreOffice profile so it never touches the user's settings."""
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run(
            ["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
             "--outdir", str(out_dir), str(docx_path)],
            check=True, capture_output=True, timeout=180,
        )
    return out_dir / f"{docx_path.stem}.pdf"


def generate(name: str, out_dir: Path) -> tuple[Path, Paper]:
    paper = build_paper(name, **SPECS[name])
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / f"{name}_lo.docx"
    build_docx(paper, docx_path)
    return to_pdf(docx_path, out_dir), paper


def main(argv: list[str]) -> int:
    out = Path(argv[1])
    names = argv[2:] or ["ieee_ml", "acm_graph"]
    for n in names:
        pdf, _ = generate(n, out)
        print(pdf)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
